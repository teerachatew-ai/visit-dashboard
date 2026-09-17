import io
import json
import os
import time
import threading
import urllib.request
from datetime import datetime, timedelta

import pandas as pd
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak,
)

app = Flask(__name__)

FONT_DIR = os.path.join(os.path.dirname(__file__), "static", "fonts")
pdfmetrics.registerFont(TTFont("Sarabun", os.path.join(FONT_DIR, "Sarabun-Regular.ttf")))
pdfmetrics.registerFont(TTFont("Sarabun-Bold", os.path.join(FONT_DIR, "Sarabun-Bold.ttf")))

SHEET_ID = "1Z6C7_4niXWeDJIvhkPzhqioKXNbxppzWSBo64MsaAw4"
SHEET_XLSX_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
TARGET_FILE = "data/target.xlsx"
STAFF_STATUS_FILE = "data/staff_status.json"
OWNER_OVERRIDES_FILE = "data/owner_overrides.json"

CACHE_SECONDS = 180
_cache = {"data": None, "ts": 0}
_lock = threading.Lock()
_staff_lock = threading.Lock()
_owner_lock = threading.Lock()

# On hosts without a persistent disk (e.g. Render's free tier), small pieces
# of durable state (staff status, owner overrides) are kept in Upstash Redis
# instead of local files, so they survive restarts/redeploys/sleep cycles.
UPSTASH_URL = os.environ.get("UPSTASH_REDIS_REST_URL", "").rstrip("/")
UPSTASH_TOKEN = os.environ.get("UPSTASH_REDIS_REST_TOKEN", "")
USE_UPSTASH = bool(UPSTASH_URL and UPSTASH_TOKEN)


def _upstash_get(key):
    req = urllib.request.Request(
        f"{UPSTASH_URL}/get/{key}",
        headers={"Authorization": f"Bearer {UPSTASH_TOKEN}"},
    )
    with urllib.request.urlopen(req, timeout=10) as resp:
        result = json.loads(resp.read().decode("utf-8"))
    return result.get("result")


def _upstash_set(key, value_str):
    req = urllib.request.Request(
        f"{UPSTASH_URL}/set/{key}",
        data=value_str.encode("utf-8"),
        headers={"Authorization": f"Bearer {UPSTASH_TOKEN}", "Content-Type": "text/plain"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=10) as resp:
        resp.read()


def _kv_get(store_key, file_path, default):
    if USE_UPSTASH:
        raw = _upstash_get(store_key)
        return json.loads(raw) if raw else default
    if not os.path.exists(file_path):
        return default
    with open(file_path, "r", encoding="utf-8") as f:
        return json.load(f)


def _kv_set(store_key, file_path, obj):
    payload = json.dumps(obj, ensure_ascii=False, indent=2)
    if USE_UPSTASH:
        _upstash_set(store_key, payload)
        return
    with open(file_path, "w", encoding="utf-8") as f:
        f.write(payload)


def load_resigned_staff():
    return set(_kv_get("staff_status", STAFF_STATUS_FILE, {"resigned": []}).get("resigned", []))


def save_resigned_staff(names):
    _kv_set("staff_status", STAFF_STATUS_FILE, {"resigned": sorted(names)})


def load_owner_overrides():
    return _kv_get("owner_overrides", OWNER_OVERRIDES_FILE, {"overrides": {}}).get("overrides", {})


def save_owner_overrides(overrides):
    _kv_set("owner_overrides", OWNER_OVERRIDES_FILE, {"overrides": overrides})

# nickname (target file) -> real name (Google Sheet visit log)
SL_NICK_TO_REAL = {
    "Tew": "Phaksiri",
    "Bas": "Teerachate",
    "O": "Prasertchai",
    "Poom": "Pisud",
    "Note": "Arthit",
    "Kid": "Somkid",
    "View": "Chatchon",
    "Aim": "Amika",
    "Big": "Ritthikorn",
}
SL_REAL_TO_NICK = {v: k for k, v in SL_NICK_TO_REAL.items()}

IE_NORMALIZE = {
    "roj": "ROJ", "rojana": "ROJ",
    "bpo": "BPO",
    "all area": "ALL AREA",
    "bhs": "BHS",
}


def normalize_ie(v):
    if v is None:
        return ""
    s = str(v).strip()
    if not s or s == "-":
        return ""
    key = s.lower()
    return IE_NORMALIZE.get(key, s.upper())


TOPIC_GROUPS = [
    ("เชิงสัมพันธ์", ["สวัสดีตามเทศกาล", "พูดคุยทั่วไป"]),
    ("สัญญา", ["สัญญาซื้อขาย", "BG", "Contract"]),
    ("ราคา/ต้นทุน", ["โครงสร้างราคา", "ส่วนต่างค่าก๊าซ"]),
    ("เทคนิค/บริการ", ["service flow", "ส่งผล PSV", "Site survey", "Survey", "Safety Training"]),
    ("ประชุม", ["Online meeting"]),
]


def group_topic(t):
    t = (t or "").strip()
    if not t:
        return "อื่นๆ"
    for group, keys in TOPIC_GROUPS:
        for k in keys:
            if k.lower() in t.lower():
                return group
    return "อื่นๆ"


def fetch_workbook_bytes():
    req = urllib.request.Request(SHEET_XLSX_URL, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=30) as resp:
        return resp.read()


def load_visit_and_docs():
    raw = fetch_workbook_bytes()
    xls = pd.ExcelFile(io.BytesIO(raw))

    visit = pd.read_excel(xls, sheet_name="Visit Record", header=0, skiprows=[1])
    visit = visit.dropna(subset=["Company_Name"]).copy()
    visit["Company_Name"] = visit["Company_Name"].astype(str).str.strip()
    visit["CustomerID"] = visit["CustomerID"].fillna("").astype(str).str.strip()
    visit["IE"] = visit["IE"].apply(normalize_ie)
    visit["Date"] = pd.to_datetime(visit["Date"], errors="coerce")
    visit["NGD_Staff_Name"] = visit["NGD_Staff_Name"].fillna("").astype(str).str.strip()
    visit["Topic"] = visit["Topic"].fillna("").astype(str).str.strip()
    visit["TopicGroup"] = visit["Topic"].apply(group_topic)
    visit["Details"] = visit["Details"].fillna("").astype(str)
    visit["Contact_person"] = visit["Contact_person"].fillna("").astype(str)
    visit = visit.dropna(subset=["Date"])

    doc = pd.read_excel(xls, sheet_name="Document Receive", header=0, skiprows=[1])
    doc = doc.dropna(subset=["Company Name"]).copy()
    doc["Company Name"] = doc["Company Name"].astype(str).str.strip()
    doc["CustomerID"] = doc["CustomerID"].fillna("").astype(str).str.strip()
    doc["IE"] = doc["IE"].apply(normalize_ie)
    doc["Date"] = pd.to_datetime(doc["Date"], errors="coerce")
    doc["NGD Staff Name"] = doc["NGD Staff Name"].fillna("").astype(str).str.strip()
    for col in ["Document 1", "Document 2", "Note", "Reciever"]:
        doc[col] = doc[col].fillna("").astype(str).str.strip()
    doc = doc.dropna(subset=["Date"])

    return visit, doc


def load_targets():
    rows = []
    xls = pd.ExcelFile(TARGET_FILE)
    for sheet in ["PNGD", "ANGD"]:
        df = pd.read_excel(xls, sheet_name=sheet, header=0)
        df = df.dropna(subset=["Project ID."]).copy()
        for _, r in df.iterrows():
            plan = str(r.get("Plan Visit", "")).strip()
            if "4 เดือน" in plan:
                per_year, months = 3, 4
            elif "6 เดือน" in plan:
                per_year, months = 2, 6
            else:
                per_year, months = 2, 6
            nick = str(r.get("SL", "")).strip()
            rows.append({
                "CustomerID": str(r.get("Project ID.", "")).strip(),
                "CompanyName": str(r.get("Customer Name", "")).strip(),
                "IE": normalize_ie(r.get("IE", "")),
                "SL_nick": nick,
                "SL_real": SL_NICK_TO_REAL.get(nick, nick),
                "CSI2025": str(r.get("CSI 2025", "")).strip(),
                "TargetPerYear": per_year,
                "TargetMonths": months,
                "BusinessUnit": sheet,
            })
    return pd.DataFrame(rows)


def build_dataset():
    visit, doc = load_visit_and_docs()
    targets = load_targets()

    now = pd.Timestamp.now()
    year_start = pd.Timestamp(year=now.year, month=1, day=1)
    day_of_year = (now - year_start).days + 1
    year_fraction = day_of_year / 365.0

    # A "touchpoint" toward target = a visit OR a document receive, since both
    # involve in-person customer contact.
    touch = pd.concat([
        visit[["CustomerID", "Date"]],
        doc[["CustomerID", "Date"]],
    ], ignore_index=True)
    touch = touch[touch["CustomerID"] != ""]

    touch_this_year = touch[touch["Date"] >= year_start]
    touch_count_year = touch_this_year.groupby("CustomerID").size()
    touch_count_total = touch.groupby("CustomerID").size()
    last_touch = touch.groupby("CustomerID")["Date"].max()

    # Topic of each customer's most recent VISIT specifically (document
    # receipts don't carry a topic), for display alongside "ติดต่อล่าสุด".
    last_visit_topic = (
        visit.sort_values("Date", ascending=False)
        .drop_duplicates(subset="CustomerID", keep="first")
        .set_index("CustomerID")["Topic"]
    )

    elapsed_months = year_fraction * 12
    owner_overrides = load_owner_overrides()

    customers = []
    for _, t in targets.iterrows():
        cid = t["CustomerID"]
        override_owner = owner_overrides.get(cid)
        sl_real = override_owner if override_owner else t["SL_real"]
        sl_nick = SL_REAL_TO_NICK.get(sl_real, "") if override_owner else t["SL_nick"]
        v_year = int(touch_count_year.get(cid, 0))
        v_total = int(touch_count_total.get(cid, 0))
        lv = last_touch.get(cid)
        target_per_year = t["TargetPerYear"]
        target_months = t["TargetMonths"]
        needed = max(0, round(target_per_year - v_year))

        days_since = None if pd.isna(lv) else (now - lv).days

        # Calendar-checkpoint based overdue logic: the year is split into
        # `target_per_year` checkpoints spaced `target_months` apart. As long
        # as the customer has met the visit count required by the most
        # recently passed checkpoint, they're not overdue - regardless of
        # exactly when within the year those visits happened.
        checkpoints_passed = int(elapsed_months // target_months)
        required_by_now = min(checkpoints_passed, target_per_year)

        if v_year >= target_per_year:
            status = "done"
        elif v_year < required_by_now:
            status = "overdue"
        else:
            status = "in_progress"

        next_checkpoint_num = v_year + 1
        next_due = (
            year_start + pd.DateOffset(months=int(target_months * next_checkpoint_num))
            if next_checkpoint_num <= target_per_year else None
        )

        customers.append({
            "customer_id": cid,
            "company_name": t["CompanyName"],
            "ie": t["IE"],
            "sl_real": sl_real,
            "sl_nick": sl_nick,
            "owner_overridden": bool(override_owner),
            "business_unit": t["BusinessUnit"],
            "csi2025": t["CSI2025"],
            "target_per_year": target_per_year,
            "target_months": target_months,
            "visits_this_year": v_year,
            "visits_total": v_total,
            "visits_needed": int(needed),
            "last_visit_date": lv.strftime("%Y-%m-%d") if pd.notna(lv) else None,
            "last_visit_topic": last_visit_topic.get(cid, "") or "",
            "days_since_last_visit": int(days_since) if days_since is not None else None,
            "next_due_date": next_due.strftime("%Y-%m-%d") if next_due is not None and pd.notna(next_due) else None,
            "status": status,
        })

    target_ids = set(targets["CustomerID"])
    off_target_ids = sorted(set(visit["CustomerID"]) - target_ids - {""})
    off_target = []
    for cid in off_target_ids:
        sub = visit[visit["CustomerID"] == cid]
        off_target.append({
            "customer_id": cid,
            "company_name": sub["Company_Name"].iloc[0],
            "ie": sub["IE"].iloc[-1],
            "visits_total": int(len(sub)),
            "visits_this_year": int(len(sub[sub["Date"] >= year_start])),
            "last_visit_date": sub["Date"].max().strftime("%Y-%m-%d"),
        })

    visits_out = []
    for _, r in visit.sort_values("Date", ascending=False).iterrows():
        visits_out.append({
            "company_name": r["Company_Name"],
            "customer_id": r["CustomerID"],
            "ie": r["IE"],
            "date": r["Date"].strftime("%Y-%m-%d"),
            "month": r["Date"].strftime("%Y-%m"),
            "topic": r["Topic"],
            "topic_group": r["TopicGroup"],
            "details": r["Details"],
            "contact": r["Contact_person"],
            "staff": r["NGD_Staff_Name"],
            "sl_nick": SL_REAL_TO_NICK.get(r["NGD_Staff_Name"], r["NGD_Staff_Name"]),
        })

    docs_out = []
    for _, r in doc.sort_values("Date", ascending=False).iterrows():
        docs_out.append({
            "company_name": r["Company Name"],
            "customer_id": r["CustomerID"],
            "ie": r["IE"],
            "date": r["Date"].strftime("%Y-%m-%d"),
            "doc1": r["Document 1"],
            "doc2": r["Document 2"],
            "note": r["Note"],
            "receiver": r["Reciever"],
            "staff": r["NGD Staff Name"],
        })

    all_staff = sorted(set(t["sl_real"] for t in customers) | set(v["staff"] for v in visits_out if v["staff"]))
    resigned = load_resigned_staff()
    sl_list = [s for s in all_staff if s not in resigned]

    return {
        "generated_at": now.strftime("%Y-%m-%d %H:%M:%S"),
        "customers": customers,
        "off_target_customers": off_target,
        "visits": visits_out,
        "documents": docs_out,
        "sl_list": sl_list,
        "all_staff": all_staff,
        "resigned_staff": sorted(resigned),
    }


def get_dataset(force=False):
    with _lock:
        if not force and _cache["data"] is not None and (time.time() - _cache["ts"]) < CACHE_SECONDS:
            return _cache["data"]
        data = build_dataset()
        _cache["data"] = data
        _cache["ts"] = time.time()
        return data


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/data")
def api_data():
    return jsonify(get_dataset())


@app.route("/api/refresh", methods=["POST"])
def api_refresh():
    data = get_dataset(force=True)
    return jsonify({"ok": True, "generated_at": data["generated_at"]})


@app.route("/api/staff/set_status", methods=["POST"])
def api_staff_set_status():
    body = request.get_json(force=True) or {}
    name = str(body.get("name", "")).strip()
    resigned = bool(body.get("resigned", False))
    if not name:
        return jsonify({"ok": False, "error": "missing name"}), 400
    with _staff_lock:
        resigned_set = load_resigned_staff()
        if resigned:
            resigned_set.add(name)
        else:
            resigned_set.discard(name)
        save_resigned_staff(resigned_set)
    get_dataset(force=True)
    return jsonify({"ok": True, "resigned_staff": sorted(resigned_set)})


@app.route("/api/owners/export")
def api_owners_export():
    data = get_dataset()
    rows = [{
        "CustomerID": c["customer_id"],
        "บริษัท": c["company_name"],
        "พื้นที่ (IE)": c["ie"],
        "Owner ปัจจุบัน": c["sl_real"],
        "Owner ใหม่ (กรอกเฉพาะรายที่ต้องการเปลี่ยน)": "",
    } for c in data["customers"]]
    df = pd.DataFrame(rows)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="โอนย้าย Owner")
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name="owner_transfer_form.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/api/owners/import", methods=["POST"])
def api_owners_import():
    file = request.files.get("file")
    if not file:
        return jsonify({"ok": False, "error": "ไม่พบไฟล์ที่อัปโหลด"}), 400

    try:
        df = pd.read_excel(file, header=0)
    except Exception as e:
        return jsonify({"ok": False, "error": f"อ่านไฟล์ Excel ไม่สำเร็จ: {e}"}), 400

    cols = {str(c).strip().lower(): c for c in df.columns}
    id_col = next((cols[k] for k in cols if "customerid" in k.replace(" ", "")), None)
    new_owner_col = next((cols[k] for k in cols if "ใหม่" in k or "new" in k), None)
    if id_col is None or new_owner_col is None:
        return jsonify({"ok": False, "error": "ไม่พบคอลัมน์ CustomerID หรือ 'Owner ใหม่' ในไฟล์ กรุณาใช้แบบฟอร์มที่ดาวน์โหลดจากระบบ"}), 400

    current_data = get_dataset()
    valid_customers = {c["customer_id"]: c for c in current_data["customers"]}

    with _owner_lock:
        overrides = load_owner_overrides()
        applied = []
        skipped = []
        for _, row in df.iterrows():
            cid = str(row.get(id_col, "")).strip()
            new_owner = str(row.get(new_owner_col, "")).strip()
            if not cid or not new_owner or new_owner.lower() == "nan":
                continue
            if cid not in valid_customers:
                skipped.append({"customer_id": cid, "reason": "ไม่พบรหัสลูกค้านี้ในระบบ"})
                continue
            old_owner = valid_customers[cid]["sl_real"]
            overrides[cid] = new_owner
            applied.append({
                "customer_id": cid,
                "company_name": valid_customers[cid]["company_name"],
                "old_owner": old_owner,
                "new_owner": new_owner,
            })
        save_owner_overrides(overrides)

    get_dataset(force=True)
    return jsonify({"ok": True, "applied": applied, "skipped": skipped, "applied_count": len(applied), "skipped_count": len(skipped)})


@app.route("/api/owners/reset", methods=["POST"])
def api_owners_reset():
    body = request.get_json(force=True) or {}
    cid = str(body.get("customer_id", "")).strip()
    if not cid:
        return jsonify({"ok": False, "error": "missing customer_id"}), 400
    with _owner_lock:
        overrides = load_owner_overrides()
        overrides.pop(cid, None)
        save_owner_overrides(overrides)
    get_dataset(force=True)
    return jsonify({"ok": True})


def build_visit_report_groups():
    """Current-year visit records grouped by each customer's CURRENT owner
    (not the staff name logged on the visit itself, which drifts as
    ownership rotates - see the ทีม tab for the same convention)."""
    data = get_dataset()
    year = str(datetime.now().year)
    owner_map = {c["customer_id"]: c["sl_real"] for c in data["customers"]}
    UNASSIGNED = "ไม่มี Owner (นอกแผนเป้าหมาย)"

    by_owner = {}
    for v in data["visits"]:
        if not v["date"].startswith(year):
            continue
        owner = owner_map.get(v["customer_id"], UNASSIGNED)
        by_owner.setdefault(owner, []).append(v)

    owners_sorted = sorted(by_owner.keys(), key=lambda o: (o == UNASSIGNED, o))
    groups = []
    for owner in owners_sorted:
        rows = sorted(by_owner[owner], key=lambda v: v["date"], reverse=True)
        groups.append((owner, rows))
    return groups, data["generated_at"], year


def build_audit_summary():
    """Figures for the audit summary page: scope, criteria, per-owner
    results, and the exception list (customers currently overdue)."""
    data = get_dataset()
    year = str(datetime.now().year)
    customers = data["customers"]
    owner_map = {c["customer_id"]: c["sl_real"] for c in customers}

    # Scope every count to in-plan customers so the per-owner rows add up to
    # the org total - an auditor will check that the columns reconcile.
    # Contacts with off-plan customers are reported separately as a note.
    all_visits_year = [v for v in data["visits"] if v["date"].startswith(year)]
    all_docs_year = [d for d in data["documents"] if d["date"].startswith(year)]
    visits_year = [v for v in all_visits_year if v["customer_id"] in owner_map]
    docs_year = [d for d in all_docs_year if d["customer_id"] in owner_map]
    off_plan_touchpoints = (len(all_visits_year) - len(visits_year)) + (len(all_docs_year) - len(docs_year))

    # Touchpoints are credited to the customer's current owner, matching how
    # the dashboard counts them everywhere else.
    touch_by_owner = {}
    for rec in visits_year + docs_year:
        owner = owner_map.get(rec["customer_id"])
        if owner:
            touch_by_owner[owner] = touch_by_owner.get(owner, 0) + 1

    per_owner = {}
    for c in customers:
        stat = per_owner.setdefault(
            c["sl_real"], {"total": 0, "done": 0, "in_progress": 0, "overdue": 0}
        )
        stat["total"] += 1
        stat[c["status"]] += 1

    rows = []
    for owner in sorted(per_owner):
        s = per_owner[owner]
        rows.append({
            "owner": owner,
            "customers": s["total"],
            "touchpoints": touch_by_owner.get(owner, 0),
            "done": s["done"],
            "in_progress": s["in_progress"],
            "overdue": s["overdue"],
            "pct_done": (s["done"] / s["total"] * 100) if s["total"] else 0,
        })

    totals = {
        "customers": len(customers),
        "visits": len(visits_year),
        "docs": len(docs_year),
        "touchpoints": len(visits_year) + len(docs_year),
        "done": sum(1 for c in customers if c["status"] == "done"),
        "in_progress": sum(1 for c in customers if c["status"] == "in_progress"),
        "overdue": sum(1 for c in customers if c["status"] == "overdue"),
    }
    totals["pct_done"] = (totals["done"] / totals["customers"] * 100) if totals["customers"] else 0

    exceptions = sorted(
        (c for c in customers if c["status"] == "overdue"),
        key=lambda c: (c["sl_real"], -(c["days_since_last_visit"] or 0)),
    )

    return {
        "year": year,
        "generated_at": data["generated_at"],
        "rows": rows,
        "totals": totals,
        "exceptions": exceptions,
        "off_plan_touchpoints": off_plan_touchpoints,
    }


def audit_criteria_lines(summary):
    """Assessment criteria spelled out for the auditor, so the numbers on
    the summary page can be traced back to a stated rule."""
    return [
        f"ขอบเขตการตรวจ: ลูกค้าในแผนเป้าหมาย {summary['totals']['customers']:,} ราย "
        f"รอบปี {summary['year']} (1 ม.ค. {summary['year']} – ปัจจุบัน)",
        "เกณฑ์เป้าหมาย: ลูกค้าแผน “ทุกๆ 4 เดือน” = 3 ครั้ง/ปี และแผน “ทุกๆ 6 เดือน” = 2 ครั้ง/ปี",
        "การนับจำนวนครั้ง: นับทั้งการเข้าเยี่ยม (Visit Record) และการรับเอกสาร (Document Receive) "
        "เป็นการติดต่อลูกค้า 1 ครั้ง",
        "เกณฑ์ “เกินกำหนด”: แบ่งปีออกเป็นช่วงตามรอบเป้าหมาย (ทุก 4 หรือ 6 เดือน) "
        "หากผ่านจุดตรวจของรอบแล้วยังติดต่อไม่ครบจำนวนที่ควรได้ ถือว่าเกินกำหนด",
        "ผู้ดูแล (Owner): อ้างอิงผู้ดูแลที่รับผิดชอบ ณ วันที่ออกรายงาน "
        "(กรณีมีการโอนย้ายระหว่างปี ประวัติการเยี่ยมเดิมจะถูกนับให้ผู้ดูแลปัจจุบัน)",
        f"แหล่งข้อมูล: Google Sheet – ชีต Visit Record และ Document Receive "
        f"ดึงข้อมูลสด ณ {summary['generated_at']}",
        f"ข้อมูลที่ไม่นับรวม: การติดต่อลูกค้านอกแผนเป้าหมาย {summary['off_plan_touchpoints']:,} ครั้ง "
        "(ไม่นำมาประเมินผล เนื่องจากอยู่นอกขอบเขตการตรวจ)",
    ]


REPORT_COLUMNS = [
    ("date", "วันที่", 12),
    ("company_name", "บริษัท", 40),
    ("ie", "พื้นที่", 10),
    ("topic", "หัวข้อ", 28),
    ("staff", "ผู้บันทึก", 16),
]


SUMMARY_HEADERS = [
    ("owner", "ผู้ดูแล (Owner)", 22),
    ("customers", "ลูกค้าที่รับผิดชอบ", 18),
    ("touchpoints", "จำนวนติดต่อปีนี้", 18),
    ("done", "ครบตามเป้าหมาย", 18),
    ("in_progress", "อยู่ระหว่างดำเนินการ", 20),
    ("overdue", "เกินกำหนด", 14),
    ("pct_done", "% ครบเป้าหมาย", 16),
]


def _write_audit_summary_sheet(ws, summary):
    title_font = Font(bold=True, size=14)
    section_font = Font(bold=True, size=11, color="1E3A8A")
    header_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(bold=True, color="FFFFFF")
    total_fill = PatternFill("solid", fgColor="EAF0FE")
    wrap = Alignment(wrap_text=True, vertical="top")

    for i, (_, _, width) in enumerate(SUMMARY_HEADERS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ncols = len(SUMMARY_HEADERS)

    def section(row, text):
        ws.cell(row=row, column=1, value=text).font = section_font
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        return row + 1

    ws.cell(row=1, column=1, value=f"รายงานสรุปการเข้าเยี่ยมลูกค้า ประจำปี {summary['year']}").font = title_font
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.cell(row=2, column=1, value="เอกสารประกอบการตรวจสอบ (Audit) – การปฏิบัติตามแผนการเข้าเยี่ยมลูกค้า").font = Font(size=10, color="6B7280")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)

    r = section(4, "1. ขอบเขตและเกณฑ์การตรวจสอบ")
    for line in audit_criteria_lines(summary):
        ws.cell(row=r, column=1, value="•  " + line).alignment = wrap
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        ws.row_dimensions[r].height = 28
        r += 1

    r = section(r + 1, "2. สรุปผลรวมทั้งองค์กร")
    t = summary["totals"]
    kpis = [
        ("ลูกค้าในแผนทั้งหมด", f"{t['customers']:,} ราย"),
        ("จำนวนการติดต่อปีนี้", f"{t['touchpoints']:,} ครั้ง (เยี่ยม {t['visits']:,} / เอกสาร {t['docs']:,})"),
        ("ครบตามเป้าหมายแล้ว", f"{t['done']:,} ราย ({t['pct_done']:.1f}%)"),
        ("อยู่ระหว่างดำเนินการ", f"{t['in_progress']:,} ราย"),
        ("เกินกำหนด (ข้อยกเว้น)", f"{t['overdue']:,} ราย"),
    ]
    for label, value in kpis:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value=value)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=ncols)
        r += 1

    r = section(r + 1, "3. ผลการปฏิบัติงานรายผู้ดูแล (Owner)")
    for c, (_, label, _) in enumerate(SUMMARY_HEADERS, start=1):
        cell = ws.cell(row=r, column=c, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap
    r += 1
    for row in summary["rows"]:
        for c, (key, _, _) in enumerate(SUMMARY_HEADERS, start=1):
            value = f"{row[key]:.1f}%" if key == "pct_done" else row[key]
            ws.cell(row=r, column=c, value=value)
        r += 1
    for c, (key, _, _) in enumerate(SUMMARY_HEADERS, start=1):
        if key == "owner":
            value = "รวมทั้งหมด"
        elif key == "pct_done":
            value = f"{t['pct_done']:.1f}%"
        else:
            value = t[key]
        cell = ws.cell(row=r, column=c, value=value)
        cell.font = Font(bold=True)
        cell.fill = total_fill
    r += 2

    r = section(r, f"4. รายการที่ไม่เป็นไปตามเกณฑ์ – เกินกำหนดเข้าเยี่ยม ({len(summary['exceptions'])} ราย)")
    if summary["exceptions"]:
        exc_headers = ["บริษัท", "พื้นที่", "ผู้ดูแล", "เป้าหมาย/ปี", "ติดต่อแล้ว", "ติดต่อล่าสุด", "ห่างมาแล้ว (วัน)"]
        for c, label in enumerate(exc_headers, start=1):
            cell = ws.cell(row=r, column=c, value=label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = wrap
        r += 1
        for c_ in summary["exceptions"]:
            values = [
                c_["company_name"], c_["ie"], c_["sl_real"], c_["target_per_year"],
                c_["visits_this_year"], c_["last_visit_date"] or "ยังไม่เคย",
                c_["days_since_last_visit"] if c_["days_since_last_visit"] is not None else "-",
            ]
            for c, value in enumerate(values, start=1):
                ws.cell(row=r, column=c, value=value).alignment = wrap
            r += 1
    else:
        ws.cell(row=r, column=1, value="ไม่พบลูกค้าที่เกินกำหนดเข้าเยี่ยม ณ วันที่ออกรายงาน")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        r += 1

    r = section(r + 1, "5. ผู้จัดทำและผู้ตรวจสอบ")
    ws.cell(row=r + 1, column=1, value="ผู้จัดทำรายงาน: ____________________")
    ws.cell(row=r + 1, column=4, value="ผู้ตรวจสอบ: ____________________")
    ws.cell(row=r + 2, column=1, value="วันที่: ____________________")
    ws.cell(row=r + 2, column=4, value="วันที่: ____________________")


@app.route("/api/report/excel")
def api_report_excel():
    groups, generated_at, year = build_visit_report_groups()
    summary = build_audit_summary()

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "สรุปสำหรับ Audit"
    _write_audit_summary_sheet(ws_summary, summary)

    ws = wb.create_sheet("รายละเอียดการเยี่ยม")
    ws.sheet_properties.outlinePr.summaryBelow = False  # group header stays above its collapsed rows

    header_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(bold=True, color="FFFFFF")
    group_fill = PatternFill("solid", fgColor="EAF0FE")
    group_font = Font(bold=True, color="1E3A8A", size=12)
    wrap = Alignment(wrap_text=True, vertical="top")

    for i, (_, label, width) in enumerate(REPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.cell(row=1, column=1, value=f"รายละเอียดการเข้าเยี่ยมลูกค้า ปี {year} — อัปเดตล่าสุด {generated_at}").font = Font(bold=True, size=13)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(REPORT_COLUMNS))
    row_num = 3

    for owner, rows in groups:
        ws.cell(row=row_num, column=1, value=f"Owner: {owner}  ({len(rows)} ครั้ง)")
        ws.cell(row=row_num, column=1).font = group_font
        for c in range(1, len(REPORT_COLUMNS) + 1):
            ws.cell(row=row_num, column=c).fill = group_fill
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=len(REPORT_COLUMNS))
        row_num += 1

        header_row = row_num
        for c, (_, label, _) in enumerate(REPORT_COLUMNS, start=1):
            cell = ws.cell(row=header_row, column=c, value=label)
            cell.font = header_font
            cell.fill = header_fill
        row_num += 1

        for v in rows:
            for c, (key, _, _) in enumerate(REPORT_COLUMNS, start=1):
                cell = ws.cell(row=row_num, column=c, value=v.get(key, ""))
                cell.alignment = wrap
            ws.row_dimensions[row_num].outlineLevel = 1
            row_num += 1

        row_num += 1  # spacer between owner groups

    ws.freeze_panes = "A4"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"visit_report_{year}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


TABLE_STYLE_BASE = [
    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
    ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ("LEFTPADDING", (0, 0), (-1, -1), 4),
    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ("TOPPADDING", (0, 0), (-1, -1), 3),
    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
]


def _audit_summary_story(summary, styles):
    """First page of the PDF: the audit-facing summary."""
    title_style, meta_style, section_style, cell_style, header_cell_style, body_style = styles
    t = summary["totals"]
    story = [
        Paragraph(f"รายงานสรุปการเข้าเยี่ยมลูกค้า ประจำปี {summary['year']}", title_style),
        Paragraph(
            "เอกสารประกอบการตรวจสอบ (Audit) – การปฏิบัติตามแผนการเข้าเยี่ยมลูกค้า"
            f" &nbsp;|&nbsp; ข้อมูล ณ {summary['generated_at']}",
            meta_style,
        ),
        Paragraph("1. ขอบเขตและเกณฑ์การตรวจสอบ", section_style),
    ]
    for line in audit_criteria_lines(summary):
        story.append(Paragraph(f"•&nbsp; {line}", body_style))

    story.append(Paragraph("2. สรุปผลรวมทั้งองค์กร", section_style))
    kpi_rows = [[
        Paragraph("ลูกค้าในแผนทั้งหมด", header_cell_style),
        Paragraph("การติดต่อปีนี้", header_cell_style),
        Paragraph("ครบตามเป้าหมาย", header_cell_style),
        Paragraph("อยู่ระหว่างดำเนินการ", header_cell_style),
        Paragraph("เกินกำหนด (ข้อยกเว้น)", header_cell_style),
    ], [
        Paragraph(f"{t['customers']:,} ราย", cell_style),
        Paragraph(f"{t['touchpoints']:,} ครั้ง<br/>(เยี่ยม {t['visits']:,} / เอกสาร {t['docs']:,})", cell_style),
        Paragraph(f"{t['done']:,} ราย ({t['pct_done']:.1f}%)", cell_style),
        Paragraph(f"{t['in_progress']:,} ราย", cell_style),
        Paragraph(f"{t['overdue']:,} ราย", cell_style),
    ]]
    kpi_table = Table(kpi_rows, colWidths=[52 * mm] * 5)
    kpi_table.setStyle(TableStyle(TABLE_STYLE_BASE))
    story.append(kpi_table)

    story.append(Paragraph("3. ผลการปฏิบัติงานรายผู้ดูแล (Owner)", section_style))
    owner_rows = [[Paragraph(label, header_cell_style) for _, label, _ in SUMMARY_HEADERS]]
    for row in summary["rows"]:
        owner_rows.append([
            Paragraph(f"{row[key]:.1f}%" if key == "pct_done" else str(row[key]), cell_style)
            for key, _, _ in SUMMARY_HEADERS
        ])
    owner_rows.append([
        Paragraph(f"<b>{'รวมทั้งหมด' if key == 'owner' else (f'{t[key]:.1f}%' if key == 'pct_done' else f'{t[key]:,}')}</b>", cell_style)
        for key, _, _ in SUMMARY_HEADERS
    ])
    owner_table = Table(owner_rows, colWidths=[44 * mm, 33 * mm, 33 * mm, 33 * mm, 38 * mm, 28 * mm, 33 * mm], repeatRows=1)
    owner_table.setStyle(TableStyle(TABLE_STYLE_BASE + [
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#EAF0FE")),
    ]))
    story.append(owner_table)

    story.append(Paragraph(
        f"4. รายการที่ไม่เป็นไปตามเกณฑ์ – เกินกำหนดเข้าเยี่ยม ({len(summary['exceptions'])} ราย)",
        section_style,
    ))
    if summary["exceptions"]:
        exc_headers = ["บริษัท", "พื้นที่", "ผู้ดูแล", "เป้าหมาย/ปี", "ติดต่อแล้ว", "ติดต่อล่าสุด", "ห่างมาแล้ว (วัน)"]
        exc_rows = [[Paragraph(h, header_cell_style) for h in exc_headers]]
        for c in summary["exceptions"]:
            exc_rows.append([
                Paragraph(c["company_name"], cell_style),
                Paragraph(c["ie"], cell_style),
                Paragraph(c["sl_real"], cell_style),
                Paragraph(str(c["target_per_year"]), cell_style),
                Paragraph(str(c["visits_this_year"]), cell_style),
                Paragraph(c["last_visit_date"] or "ยังไม่เคย", cell_style),
                Paragraph(str(c["days_since_last_visit"] if c["days_since_last_visit"] is not None else "-"), cell_style),
            ])
        exc_table = Table(exc_rows, colWidths=[85 * mm, 20 * mm, 30 * mm, 25 * mm, 25 * mm, 30 * mm, 27 * mm], repeatRows=1)
        exc_table.setStyle(TableStyle(TABLE_STYLE_BASE))
        story.append(exc_table)
    else:
        story.append(Paragraph("ไม่พบลูกค้าที่เกินกำหนดเข้าเยี่ยม ณ วันที่ออกรายงาน", body_style))

    story.append(Paragraph("5. ผู้จัดทำและผู้ตรวจสอบ", section_style))
    sign_rows = [[
        Paragraph("ผู้จัดทำรายงาน<br/><br/>ลงชื่อ ______________________<br/>วันที่ ______________________", cell_style),
        Paragraph("ผู้ตรวจสอบ<br/><br/>ลงชื่อ ______________________<br/>วันที่ ______________________", cell_style),
    ]]
    sign_table = Table(sign_rows, colWidths=[90 * mm, 90 * mm])
    sign_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(sign_table)
    story.append(PageBreak())
    return story


@app.route("/api/report/pdf")
def api_report_pdf():
    groups, generated_at, year = build_visit_report_groups()
    summary = build_audit_summary()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        topMargin=14 * mm, bottomMargin=14 * mm, leftMargin=12 * mm, rightMargin=12 * mm,
    )

    title_style = ParagraphStyle("Title", fontName="Sarabun-Bold", fontSize=16, spaceAfter=4)
    meta_style = ParagraphStyle("Meta", fontName="Sarabun", fontSize=9, textColor=colors.grey, spaceAfter=10)
    owner_style = ParagraphStyle("Owner", fontName="Sarabun-Bold", fontSize=12, textColor=colors.HexColor("#1E3A8A"), spaceBefore=12, spaceAfter=6)
    section_style = ParagraphStyle("Section", fontName="Sarabun-Bold", fontSize=11, textColor=colors.HexColor("#1E3A8A"), spaceBefore=12, spaceAfter=6)
    body_style = ParagraphStyle("Body", fontName="Sarabun", fontSize=9, leading=13, spaceAfter=2)
    cell_style = ParagraphStyle("Cell", fontName="Sarabun", fontSize=8, leading=10)
    header_cell_style = ParagraphStyle("HeaderCell", fontName="Sarabun-Bold", fontSize=8.5, textColor=colors.white, leading=10)

    story = _audit_summary_story(
        summary,
        (title_style, meta_style, section_style, cell_style, header_cell_style, body_style),
    )
    story += [
        Paragraph(f"รายละเอียดการเข้าเยี่ยมลูกค้า ปี {year}", title_style),
        Paragraph(f"อัปเดตล่าสุด {generated_at}", meta_style),
    ]

    col_widths = [22 * mm, 95 * mm, 20 * mm, 95 * mm, 30 * mm]

    for owner, rows in groups:
        story.append(Paragraph(f"Owner: {owner} &nbsp;&nbsp;({len(rows)} ครั้ง)", owner_style))

        table_rows = [[Paragraph(label, header_cell_style) for _, label, _ in REPORT_COLUMNS]]
        for v in rows:
            table_rows.append([
                Paragraph(str(v.get(key, "")), cell_style) for key, _, _ in REPORT_COLUMNS
            ])

        table = Table(table_rows, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle(TABLE_STYLE_BASE))
        story.append(table)

    doc.build(story)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"visit_report_{year}.pdf",
        mimetype="application/pdf",
    )


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5050))
    app.run(host="0.0.0.0", port=port, debug=not os.environ.get("PORT"))
